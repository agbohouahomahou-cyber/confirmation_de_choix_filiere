import csv
from io import BytesIO
from xml.sax.saxutils import escape

from django.contrib.auth.decorators import login_required, permission_required
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from declaration.models import Decision, Ecole, Financement

from .views import PERMISSION, candidatures_filtrees


# ---------------------------------------------------------------- données

def _date(valeur):
    return valeur.strftime('%d/%m/%Y') if valeur else ''


def _recu_le(c):
    return timezone.localtime(c.cree_le).strftime('%d/%m/%Y %H:%M')


def _handicap(c):
    texte = c.get_handicap_display()
    if c.handicap_precision:
        texte += f' ({c.handicap_precision})'
    return texte


# (titre de colonne, fonction qui lit la valeur sur une candidature)
COLONNES = [
    ('NPI', lambda c: c.npi),
    ('N° de table du BAC', lambda c: c.numero_table),
    ('Nom', lambda c: c.nom),
    ('Prénoms', lambda c: c.prenoms),
    ('Date de naissance', lambda c: _date(c.date_naissance)),
    ('Lieu de naissance', lambda c: c.lieu_naissance),
    ('Téléphone', lambda c: c.telephone),
    ('Handicap', _handicap),
    ('Père ou tuteur', lambda c: f'{c.pere_nom} {c.pere_prenoms}'.strip()),
    ('Mère ou tutrice', lambda c: f'{c.mere_nom} {c.mere_prenoms}'.strip()),
    ('Décision', lambda c: c.get_decision_display()),
    ('École', lambda c: c.ecole),
    ('Financement', lambda c: c.get_financement_display()),
    ('Reçue le', _recu_le),
]


def neutralise(valeur):
    """Empêche Excel d'interpréter un texte saisi par un candidat comme une formule."""
    if isinstance(valeur, str) and valeur:
        if valeur[0] in '=@':
            return "'" + valeur
        if valeur[0] in '+-' and not valeur[1:].replace(' ', '').isdigit():
            return "'" + valeur  # un numéro comme +229 97... reste intact
    return valeur


def lignes(candidatures):
    for c in candidatures.iterator():
        yield [neutralise(lire(c)) for _, lire in COLONNES]


def description_filtres(filtres):
    morceaux = []
    if filtres['q']:
        morceaux.append(f'Recherche : « {filtres["q"]} »')
    if filtres['ecole']:
        morceaux.append(f'École : {dict(Ecole.choices)[filtres["ecole"]]}')
    if filtres['decision']:
        morceaux.append(f'Décision : {dict(Decision.choices)[filtres["decision"]]}')
    if filtres['financement']:
        morceaux.append(f'Financement : {dict(Financement.choices)[filtres["financement"]]}')
    return ' · '.join(morceaux) or 'Aucun filtre (toutes les déclarations)'


def nom_fichier(extension):
    return f'declarations_una_{timezone.localtime():%Y%m%d_%H%M}.{extension}'


def _telechargement(reponse, extension):
    reponse['Content-Disposition'] = f'attachment; filename="{nom_fichier(extension)}"'
    return reponse


# ------------------------------------------------------------------ vues

@login_required
@permission_required(PERMISSION, raise_exception=True)
def page_exports(request):
    candidatures, filtres = candidatures_filtrees(request.GET)
    contexte = {
        'filtres': filtres,
        'filtres_actifs': any(filtres.values()),
        'ecoles': Ecole.choices,
        'decisions': Decision.choices,
        'financements': Financement.choices,
    }
    return render(request, 'scolarite/exports.html', contexte)


@login_required
@permission_required(PERMISSION, raise_exception=True)
def export_csv(request):
    candidatures, _ = candidatures_filtrees(request.GET)
    reponse = HttpResponse(content_type='text/csv; charset=utf-8')
    reponse.write('\ufeff')  # BOM : Excel affiche correctement les accents
    ecrivain = csv.writer(reponse, delimiter=';')  # ';' : Excel en version française
    ecrivain.writerow([titre for titre, _ in COLONNES])
    for ligne in lignes(candidatures):
        ecrivain.writerow(ligne)
    return _telechargement(reponse, 'csv')


@login_required
@permission_required(PERMISSION, raise_exception=True)
def export_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    candidatures, filtres = candidatures_filtrees(request.GET)

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = 'Déclarations'
    feuille.append([titre for titre, _ in COLONNES])
    for ligne in lignes(candidatures):
        feuille.append(ligne)

    for cellule in feuille[1]:
        cellule.font = Font(bold=True, color='FFFFFF')
        cellule.fill = PatternFill('solid', start_color='14532D')
        cellule.alignment = Alignment(vertical='center', wrap_text=True)
    feuille.freeze_panes = 'A2'
    feuille.auto_filter.ref = feuille.dimensions
    for i in range(1, len(COLONNES) + 1):
        lettre = get_column_letter(i)
        plus_long = max(len(str(cellule.value or '')) for cellule in feuille[lettre])
        feuille.column_dimensions[lettre].width = min(max(plus_long + 2, 10), 45)

    infos = classeur.create_sheet('Informations')
    infos.append(['Export', 'Déclarations de confirmation – UNA'])
    infos.append(['Date de l\'export', timezone.localtime().strftime('%d/%m/%Y %H:%M')])
    infos.append(['Filtres', description_filtres(filtres)])
    infos.append(['Nombre de lignes', feuille.max_row - 1])
    infos.column_dimensions['A'].width = 20
    infos.column_dimensions['B'].width = 70

    reponse = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    classeur.save(reponse)
    return _telechargement(reponse, 'xlsx')


@login_required
@permission_required(PERMISSION, raise_exception=True)
def export_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    candidatures, filtres = candidatures_filtrees(request.GET)
    vert = colors.HexColor('#14532D')
    pale = colors.HexColor('#F2F5F0')

    styles = getSampleStyleSheet()
    titre = ParagraphStyle('titre', parent=styles['Title'], alignment=0, fontSize=16, textColor=vert)
    meta = ParagraphStyle('meta', parent=styles['Normal'], fontSize=9, leading=13,
                          textColor=colors.HexColor('#555555'))
    cellule = ParagraphStyle('cellule', parent=styles['Normal'], fontSize=8, leading=10)
    entete = ParagraphStyle('entete', parent=cellule, textColor=colors.white, fontName='Helvetica-Bold')

    colonnes = [
        ('Candidat', 70, lambda c: f'{c.nom} {c.prenoms}'),
        ('NPI', 30, lambda c: c.npi),
        ('École', 26, lambda c: c.ecole),
        ('Décision', 26, lambda c: 'Confirmée' if c.decision == Decision.CONFIRMER else 'Démission'),
        ('Financement', 50, lambda c: c.get_financement_display()),
        ('Téléphone', 33, lambda c: c.telephone),
        ('Reçue le', 32, _recu_le),
    ]

    total = candidatures.count()
    donnees = [[Paragraph(nom, entete) for nom, _, _ in colonnes]]
    for c in candidatures.iterator():
        donnees.append([Paragraph(escape(str(lire(c))), cellule) for _, _, lire in colonnes])

    tableau = Table(donnees, colWidths=[l * mm for _, l, _ in colonnes], repeatRows=1)
    tableau.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), vert),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, pale]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LINEBELOW', (0, 0), (-1, -1), 0.25, colors.HexColor('#D8DED5')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))

    def pied_de_page(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#777777'))
        canvas.drawString(15 * mm, 8 * mm, "Université Nationale d'Agriculture – Scolarité centrale")
        canvas.drawRightString(doc.pagesize[0] - 15 * mm, 8 * mm, f'Page {doc.page}')
        canvas.restoreState()

    tampon = BytesIO()
    document = SimpleDocTemplate(
        tampon, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=16 * mm, title='Déclarations de confirmation – UNA')
    document.build([
        Paragraph("Université Nationale d'Agriculture", meta),
        Paragraph('Déclarations de confirmation de choix', titre),
        Paragraph(f'Exporté le {timezone.localtime():%d/%m/%Y à %H:%M} · '
                  f'{total} déclaration{"s" if total > 1 else ""}', meta),
        Paragraph(escape(description_filtres(filtres)), meta),
        Spacer(1, 6 * mm),
        tableau,
    ], onFirstPage=pied_de_page, onLaterPages=pied_de_page)

    reponse = HttpResponse(tampon.getvalue(), content_type='application/pdf')
    return _telechargement(reponse, 'pdf')
